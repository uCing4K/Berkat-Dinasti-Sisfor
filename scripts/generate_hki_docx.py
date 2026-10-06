import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pygments.lexers import HtmlPhpLexer
from pygments.token import Token

def hex_to_rgb(hex_str):
    hex_str = hex_str.lstrip('#')
    return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))

# THEME DEFINITIONS
LIGHT_THEME = {
    'name': 'Modern Light',
    'bg': 'F8FAFC',
    'border': 'CBD5E1',
    'hdr_bg': '3E2B21',
    'hdr_txt': hex_to_rgb('FFFFFF'),
    'meta_bg': 'FFF7ED',
    'meta_txt': hex_to_rgb('7C2D12'),
    'default': (hex_to_rgb('1E293B'), False, False),
    'tag': (hex_to_rgb('0550AE'), True, False),         # Royal Blue Bold
    'attr': (hex_to_rgb('6F42C1'), False, False),       # Vibrant Purple
    'str': (hex_to_rgb('0A3069'), False, False),        # Navy Blue
    'kw': (hex_to_rgb('CF222E'), True, False),          # Crimson Bold
    'comment': (hex_to_rgb('6E7781'), False, True),     # Slate Grey Italic
    'func': (hex_to_rgb('8250DF'), False, False),       # Violet
    'var': (hex_to_rgb('953800'), False, False),        # Dark Amber
    'num': (hex_to_rgb('0969DA'), False, False),        # Blue
    'op': (hex_to_rgb('57606A'), False, False),         # Slate
    'punc': (hex_to_rgb('57606A'), False, False),       # Slate
    'preproc': (hex_to_rgb('D73A49'), True, False)      # Red Bold
}

DARK_THEME = {
    'name': 'VS Code Dark',
    'bg': '1E1E1E',
    'border': '334155',
    'hdr_bg': '2B1D15',
    'hdr_txt': hex_to_rgb('FFFFFF'),
    'meta_bg': '261E1A',
    'meta_txt': hex_to_rgb('FDBA74'),
    'default': (hex_to_rgb('D4D4D4'), False, False),
    'tag': (hex_to_rgb('569CD6'), True, False),         # Blue Bold
    'attr': (hex_to_rgb('9CDCFE'), False, False),       # Sky Blue
    'str': (hex_to_rgb('CE9178'), False, False),        # Warm Coral
    'kw': (hex_to_rgb('C586C0'), True, False),          # Purple Bold
    'comment': (hex_to_rgb('6A9955'), False, True),     # Olive Green Italic
    'func': (hex_to_rgb('DCDCAA'), False, False),       # Soft Yellow
    'var': (hex_to_rgb('9CDCFE'), False, False),        # Light Cyan
    'num': (hex_to_rgb('B5CEA8'), False, False),        # Mint Green
    'op': (hex_to_rgb('D4D4D4'), False, False),         # Light Silver
    'punc': (hex_to_rgb('808080'), False, False),       # Medium Silver
    'preproc': (hex_to_rgb('569CD6'), True, False)      # Blue Bold
}

def resolve_token_style(ttype, theme):
    if ttype in Token.Comment.Preproc:
        return theme['preproc']
    if ttype in Token.Comment:
        return theme['comment']
    if ttype in Token.Literal.String:
        return theme['str']
    if ttype in Token.Literal.Number or ttype in Token.Literal.Color:
        return theme['num']
    if ttype in Token.Keyword:
        return theme['kw']
    if ttype in Token.Name.Tag:
        return theme['tag']
    if ttype in Token.Name.Attribute:
        return theme['attr']
    if ttype in Token.Name.Function or ttype in Token.Name.Builtin or ttype in Token.Name.Class:
        return theme['func']
    if ttype in Token.Name.Variable:
        return theme['var']
    if ttype in Token.Operator:
        return theme['op']
    if ttype in Token.Punctuation:
        return theme['punc']
    return theme['default']

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="6", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:insideV w:val="none"/>'
                        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'</w:tblBorders>')
    tblPr.append(borders)

def make_summary_table(doc):
    headers = ["No", "Nama Berkas", "Direktori (Path)", "Bahasa", "Peran & Deskripsi Fungsional"]
    data = [
        ["1", "index.php", "/index.php", "PHP", "Entry Point utama & Router Redirect otomatis ke antarmuka login"],
        ["2", "login.php", "/public/login.php", "PHP, HTML5, Tailwind CSS, JS", "Antarmuka Autentikasi Pengguna & Branding Visual Bakery Berkat Dinasti"],
        ["3", "reset-password.php", "/public/reset-password.php", "PHP, HTML5, Tailwind CSS", "Antarmuka Alur Pemulihan Password Admin Internal beserta Info Kebijakan"]
    ]

    col_widths = [Inches(0.5), Inches(1.3), Inches(1.4), Inches(1.2), Inches(1.87)]

    table = doc.add_table(rows=len(data) + 1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="D1D5DB", sz="6")

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "3E2B21")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        r.font.name = "Segoe UI"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        bg_color = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            row_cells[col_idx].width = col_widths[col_idx]
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=100, bottom=100, left=120, right=120)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text)
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.5)
            if col_idx == 1:
                r.font.bold = True

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(6)
    sp.paragraph_format.space_after = Pt(12)

def make_highlighted_code_table(doc, title, meta_info, file_path, theme):
    with open(file_path, "r", encoding="utf-8") as f:
        code_content = f.read()

    table = doc.add_table(rows=3, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color=theme['border'], sz="6")

    for row in table.rows:
        row.cells[0].width = Inches(6.27)

    # Row 0: Title Header
    cell_0 = table.cell(0, 0)
    set_cell_background(cell_0, theme['hdr_bg'])
    set_cell_margins(cell_0, top=140, bottom=140, left=180, right=180)
    p0 = cell_0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run(title)
    r0.font.name = "Segoe UI"
    r0.font.size = Pt(10.5)
    r0.font.bold = True
    r0.font.color.rgb = theme['hdr_txt']

    # tblHeader for repetition
    trPr0 = table.rows[0]._tr.get_or_add_trPr()
    trPr0.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    # Row 1: Meta info
    cell_1 = table.cell(1, 0)
    set_cell_background(cell_1, theme['meta_bg'])
    set_cell_margins(cell_1, top=100, bottom=100, left=180, right=180)
    p1 = cell_1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run(meta_info)
    r1.font.name = "Segoe UI"
    r1.font.size = Pt(9.0)
    r1.font.color.rgb = theme['meta_txt']

    # Row 2: Code Content with Syntax Highlighting
    cell_2 = table.cell(2, 0)
    set_cell_background(cell_2, theme['bg'])
    set_cell_margins(cell_2, top=140, bottom=140, left=180, right=180)

    # Tokenize with Pygments
    lexer = HtmlPhpLexer()
    tokens = lexer.get_tokens(code_content)

    # Split into lines
    lines = [[]]
    for ttype, tvalue in tokens:
        parts = tvalue.split('\n')
        for i, part in enumerate(parts):
            if i > 0:
                lines.append([])
            if part:
                lines[-1].append((ttype, part))

    p_first = cell_2.paragraphs[0]
    p_first.paragraph_format.space_before = Pt(0)
    p_first.paragraph_format.space_after = Pt(0)
    p_first.paragraph_format.line_spacing = 1.05

    for line_idx, line_tokens in enumerate(lines):
        p = p_first if line_idx == 0 else cell_2.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05

        if not line_tokens:
            r = p.add_run(" ")
            r.font.name = "Consolas"
            r.font.size = Pt(8.5)
            continue

        for ttype, val in line_tokens:
            color, is_bold, is_italic = resolve_token_style(ttype, theme)
            r = p.add_run(val)
            r.font.name = "Consolas"
            r.font.size = Pt(8.5)
            r.font.color.rgb = color
            r.font.bold = is_bold
            r.font.italic = is_italic

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(6)
    sp.paragraph_format.space_after = Pt(12)

def generate_full_document(base_dir, output_path, theme):
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(10)
    r_title = p_title.add_run("PERANCANGAN USER INTERFACE / USER EXPERIENCE (UI/UX) BERKAT DINASTI UNTUK MENDUKUNG PENGELOLAAN PESANAN DAN INFORMASI OPERASIONAL UMKM")
    r_title.font.name = 'Segoe UI'
    r_title.font.size = Pt(14)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x3E, 0x2B, 0x21)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    r_sub = p_sub.add_run("DOKUMEN SPESIFIKASI DAN KODE SUMBER PROGRAM KOMPUTER\nPENGAJUAN HAK KEKAYAAN INTELEKTUAL (HKI)")
    r_sub.font.name = 'Segoe UI'
    r_sub.font.size = Pt(9.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(0xEA, 0x58, 0x0C)

    def add_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Segoe UI'
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x3E, 0x2B, 0x21)
        return p

    def add_subheading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Segoe UI'
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        return p

    def add_body_p(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        return p

    def add_bullet_p(num_text, text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r_num = p.add_run(num_text)
        r_num.font.name = 'Times New Roman'
        r_num.font.size = Pt(11)
        r_num.font.bold = True
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        return p

    # --- BAB I. JUDUL ---
    add_heading("I. Judul Ciptaan")
    add_body_p("PERANCANGAN USER INTERFACE / USER EXPERIENCE (UI/UX) BERKAT DINASTI UNTUK MENDUKUNG PENGELOLAAN PESANAN DAN INFORMASI OPERASIONAL UMKM")

    # --- BAB II. LATAR BELAKANG ---
    add_heading("II. Latar Belakang")
    add_body_p("Berkat Dinasti merupakan UMKM produsen dan penjual roti berbasis pre-order yang berlokasi di Desa Songokerto, Kota Batu. Berdasarkan hasil identifikasi kebutuhan pada tahap awal pengembangan, proses penerimaan pesanan dilakukan melalui WhatsApp dan tatap muka. Sementara itu, pencatatan pesanan, pembayaran, serta rekap informasi penjualan sebelumnya masih banyak dilakukan secara manual. Kondisi tersebut menimbulkan kebutuhan akan suatu rancangan antarmuka digital yang dapat membantu admin atau pemilik dalam mengelola informasi operasional secara lebih terstruktur, sederhana, dan mudah digunakan.")
    add_body_p("Berdasarkan kebutuhan tersebut, dikembangkan rancangan User Interface/User Experience (UI/UX) yang berfokus pada aktivitas internal admin atau pemilik. Perancangan tidak ditujukan untuk menggantikan mekanisme pemesanan pelanggan yang telah berjalan, karena pelanggan tetap melakukan pemesanan melalui WhatsApp maupun secara langsung. Oleh karena itu, rancangan UI/UX diarahkan untuk mendukung aktivitas internal, khususnya dalam melihat dan menambahkan pesanan, memantau jadwal atau deadline pengiriman, memperbarui status pesanan, melihat informasi produk, serta memantau informasi pemasukan.")
    add_body_p("Dalam proses pengembangan sebelumnya terdapat rancangan basis data untuk mendukung sistem. Namun, basis data tersebut pada tahap pengembangan saat ini sudah tidak digunakan dan tidak menjadi bagian utama dari karya yang diajukan. Fokus karya diarahkan sepenuhnya pada hasil perancangan UI/UX, meliputi struktur informasi, arsitektur navigasi, alur interaksi pengguna, wireframe, prototype low-fidelity dan high-fidelity, komponen antarmuka, serta keputusan desain yang dibangun berdasarkan kebutuhan pengguna.")
    add_body_p("Perancangan UI/UX dilakukan dengan pendekatan Design Thinking melalui tahapan identifikasi kebutuhan pengguna, perumusan permasalahan, pengembangan ide, pembuatan prototype, dan pengujian terhadap rancangan. Hasil pengujian digunakan untuk mengevaluasi kemudahan penggunaan dan kesesuaian rancangan dengan kebutuhan admin atau pemilik.")
    add_body_p("Perancangan ini juga memiliki keterkaitan dengan Sustainable Development Goals (SDGs), khususnya SDG 8: Decent Work and Economic Growth, melalui upaya mendukung produktivitas dan penguatan aktivitas operasional UMKM dengan pemanfaatan rancangan teknologi digital yang sesuai dengan kebutuhan pengguna. Karya ini juga mendukung SDG 9: Industry, Innovation and Infrastructure melalui pengembangan inovasi digital dalam bentuk rancangan UI/UX yang dirancang secara sistematis berdasarkan kebutuhan pengguna.")
    add_body_p("Dengan demikian, karya ini menghasilkan suatu rancangan UI/UX yang menitikberatkan pada kemudahan akses informasi, kesederhanaan navigasi, serta efektivitas pengelolaan pesanan dan informasi operasional. Hasil rancangan tersebut menjadi objek utama karya yang diajukan dalam pengurusan Hak Kekayaan Intelektual (HKI), sekaligus menjadi bentuk kontribusi pada pemanfaatan inovasi digital untuk mendukung penguatan operasional UMKM.")

    # --- BAB III. TUJUAN ---
    add_heading("III. Tujuan")
    add_bullet_p("1. ", "Menghasilkan rancangan UI/UX yang sesuai dengan kebutuhan admin/pemilik Berkat Dinasti dalam mendukung pengelolaan pesanan dan informasi operasional UMKM.")
    add_bullet_p("2. ", "Merancang struktur informasi dan alur interaksi yang sederhana, jelas, dan mudah dipahami sehingga pengguna dapat mengakses informasi operasional secara lebih efektif.")
    add_bullet_p("3. ", "Menyediakan prototype UI/UX dalam bentuk rancangan low-fidelity dan high-fidelity sebagai representasi visual dan interaktif dari rancangan yang dikembangkan.")
    add_bullet_p("4. ", "Mempermudah pengelolaan informasi pesanan, khususnya dalam aktivitas menambahkan pesanan, melihat daftar pesanan, memantau deadline pengiriman, serta memperbarui status pesanan menjadi Proses atau Selesai.")
    add_bullet_p("5. ", "Mempermudah pemantauan informasi operasional, meliputi data jenis roti/produk, informasi pelanggan, serta informasi pemasukan yang dibutuhkan oleh admin/pemilik.")
    add_bullet_p("6. ", "Menerapkan pendekatan Design Thinking dalam proses perancangan UI/UX dengan mempertimbangkan kebutuhan, permasalahan, dan hasil pengujian dari pengguna.")
    add_bullet_p("7. ", "Mendukung proses digitalisasi internal UMKM melalui rancangan antarmuka yang sederhana, informatif, dan sesuai dengan aktivitas operasional Berkat Dinasti, tanpa mengubah mekanisme pemesanan pelanggan yang tetap dilakukan melalui WhatsApp maupun secara offline.")
    add_bullet_p("8. ", "Mendukung SDG 8 (Decent Work and Economic Growth) melalui perancangan UI/UX yang membantu meningkatkan efektivitas pengelolaan informasi dan aktivitas operasional UMKM sebagai bagian dari penguatan produktivitas usaha.")
    add_bullet_p("9. ", "Mendukung SDG 9 (Industry, Innovation and Infrastructure) melalui pengembangan inovasi digital berupa rancangan UI/UX yang berorientasi pada kebutuhan pengguna dan dapat menjadi dasar pengembangan layanan digital bagi aktivitas operasional UMKM.")

    # --- BAB IV. RANCANGAN UI/UX YANG DIKEMBANGKAN ---
    add_heading("IV. Rancangan UI/UX yang Dikembangkan")
    add_body_p("Rancangan antarmuka pengguna (User Interface) dan pengalaman pengguna (User Experience) Berkat Dinasti dirancang dengan berorientasi pada kemudahan penggunaan (usability), efisiensi operasional internal UMKM, dan konsistensi visual identitas usaha bakery berbasis pre-order. Komponen antarmuka yang dikembangkan dan diwujudkan ke dalam implementasi program berbasis web meliputi:")

    add_subheading("4.1 Identitas Visual dan Sistem Desain (Design System)")
    add_body_p("Sistem desain dibangun untuk merefleksikan karakter usaha UMKM toko roti dan pastry modern dengan palet warna yang hangat, ramah, dan profesional:")
    add_bullet_p("• ", "Brand Brown (#3E2B21): Warna cokelat tua kopi/roti panggang yang memberikan kesan elegan, kokoh, dan berkarakter, digunakan pada panel branding visual utama.")
    add_bullet_p("• ", "Brand Orange (#FF9B45 & #E88C3D): Aksen warna oranye hangat khas bakery yang membangkitkan kehangatan dan nafsu makan, difungsikan sebagai warna fokus, tombol aksi utama (Call to Action), dan badge kategori.")
    add_bullet_p("• ", "Soft Background (#F9F7F5): Latar belakang krem/netral yang bersih, memberikan kenyamanan visual bagi pengguna saat bekerja dalam durasi panjang.")
    add_bullet_p("• ", "Tipografi Modern: Menggunakan jenis huruf sans-serif 'Inter' dari Google Fonts yang memiliki keterbacaan tinggi di berbagai ukuran resolusi layar perangkat.")

    add_subheading("4.2 Rancangan Antarmuka Halaman Login (public/login.php)")
    add_body_p("Halaman Login merupakan pintu gerbang keamanan utama bagi pengguna (admin/pemilik UMKM) untuk mengakses sistem operasional internal. Antarmuka ini mengusung tata letak 'Split-Screen Layout' yang responsif:")
    add_bullet_p("• ", "Panel Branding (Sisi Kiri Desktop): Menampilkan identitas visual logo Berkat Dinasti dengan efek glow lembut, tagline operasional 'Kelola Pesanan Rotimu dengan Mudah', badge penanda sistem manajemen pesanan, dan verifikasi internal admin.")
    add_bullet_p("• ", "Panel Interaksi Form (Sisi Kanan): Menampilkan formulir login yang bersih dan intuitif, dilengkapi ikon bantu pada input username dan password, tombol interaktif untuk toggle visibilitas sandi (show/hide password), opsi 'Ingat Saya' (remember me), serta tautan menuju pemulihan sandi.")

    add_subheading("4.3 Rancangan Antarmuka Halaman Reset Password (public/reset-password.php)")
    add_body_p("Halaman Reset Password dirancang untuk memfasilitasi alur penanganan masalah lupa kata sandi bagi staf/admin secara terkendali. Fitur antarmuka ini meliputi:")
    add_bullet_p("• ", "Tautan Navigasi Kembali ('Kembali ke Login'): Memudahkan pengguna kembali ke halaman utama jika membatalkan proses reset.")
    add_bullet_p("• ", "Formulir Permintaan Pemulihan: Input verifikasi username pengguna yang terintegrasi secara ringkas.")
    add_bullet_p("• ", "Notifikasi Operasional: Menyediakan kotak informasi status penanganan (SLA 1x24 jam oleh Super Admin) yang memberikan kejelasan proses bagi pengguna.")

    add_subheading("4.4 Modul Pengarah Navigasi Sistem (index.php)")
    add_body_p("Modul entry point sistem bertindak sebagai pengarah jalur otomatis (HTTP redirector) yang memastikan seluruh pengunjung yang mengakses domain/root direktori langsung diarahkan secara aman ke gerbang antarmuka login.")

    # --- BAB V. SOURCE CODE ---
    add_heading("V. Kode Sumber (Source Code) Program")
    add_body_p(f"Berikut adalah kode sumber lengkap implementasi program komputer dari antarmuka pengguna (UI/UX) Sistem Informasi Berkat Dinasti dengan penyorotan sintaks berwarna (Syntax Highlighting - Tema {theme['name']}). Seluruh kode disajikan secara transparan dan terstruktur di dalam tabel lampiran di bawah ini guna memenuhi persyaratan dokumen teknis pendaftaran Hak Cipta Program Komputer / Hak Kekayaan Intelektual (HKI):")

    add_subheading("Ringkasan Berkas Kode Sumber Antarmuka")
    make_summary_table(doc)

    # Table 1: index.php
    add_subheading("5.1 Kode Sumber File: index.php (Root Entry Point & Routing Redirector)")
    index_file = os.path.join(base_dir, "index.php")
    make_highlighted_code_table(
        doc,
        title="Berkas: index.php",
        meta_info="Lokasi: /index.php | Bahasa: PHP | Fungsi: Entry Point Sistem & Otomatisasi Pengalihan ke Antarmuka Login",
        file_path=index_file,
        theme=theme
    )

    # Table 2: public/login.php
    add_subheading("5.2 Kode Sumber File: public/login.php (Antarmuka Halaman Login Pengguna)")
    login_file = os.path.join(base_dir, "public", "login.php")
    make_highlighted_code_table(
        doc,
        title="Berkas: public/login.php",
        meta_info="Lokasi: /public/login.php | Bahasa: PHP, HTML5, Tailwind CSS, JavaScript | Fungsi: Antarmuka Autentikasi Pengguna & Branding Visual Bakery",
        file_path=login_file,
        theme=theme
    )

    # Table 3: public/reset-password.php
    add_subheading("5.3 Kode Sumber File: public/reset-password.php (Antarmuka Pemulihan Kata Sandi)")
    reset_file = os.path.join(base_dir, "public", "reset-password.php")
    make_highlighted_code_table(
        doc,
        title="Berkas: public/reset-password.php",
        meta_info="Lokasi: /public/reset-password.php | Bahasa: PHP, HTML5, Tailwind CSS | Fungsi: Antarmuka Alur Pemulihan Password Admin Internal",
        file_path=reset_file,
        theme=theme
    )

    doc.save(output_path)
    print(f"Berhasil menghasilkan dokumen HKI DOCX ({theme['name']}): {output_path}")

def update_markdown_file(base_dir, md_path):
    with open(os.path.join(base_dir, "index.php"), "r", encoding="utf-8") as f:
        index_code = f.read().strip()

    with open(os.path.join(base_dir, "public", "login.php"), "r", encoding="utf-8") as f:
        login_code = f.read().strip()

    with open(os.path.join(base_dir, "public", "reset-password.php"), "r", encoding="utf-8") as f:
        reset_code = f.read().strip()

    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    marker = "IV. Rancangan UI/UX yang Dikembangkan"
    idx = content.find(marker)
    if idx == -1:
        marker = "IV. Rancangan UI/UX yang dikembangkan"
        idx = content.find(marker)
    
    if idx != -1:
        header_part = content[:idx].strip()
    else:
        header_part = content.strip()

    new_content = header_part + """

IV. Rancangan UI/UX yang Dikembangkan

Rancangan antarmuka pengguna (*User Interface*) dan pengalaman pengguna (*User Experience*) Berkat Dinasti dirancang dengan menitikberatkan pada kemudahan penggunaan (*usability*), efisiensi alur kerja operasional internal UMKM, dan konsistensi visual identitas usaha bakery berbasis *pre-order*. Komponen antarmuka yang dikembangkan dan diwujudkan ke dalam implementasi program komputer berbasis web meliputi:

1. **Identitas Visual dan Sistem Desain (*Design System*)**
   - **Palet Warna Tematik Bakery**: Menggunakan kombinasi warna *Brand Brown* (`#3E2B21`) untuk panel identitas utama, *Brand Orange* (`#FF9B45` & `#E88C3D`) sebagai warna aksen interaktif dan tombol aksi utama (*Call to Action*), serta *Soft Neutral Background* (`#F9F7F5`) sebagai latar belakang kerja yang bersih dan ramah di mata.
   - **Tipografi Modern**: Menggunakan jenis huruf sans-serif modern *Inter* dari Google Fonts untuk memastikan keterbacaan (*readability*) optimal pada layar komputer desktop maupun perangkat bergerak.
   - **Tata Letak Responsif (*Split-Screen Layout*)**: Memisahkan area *branding display* informatif di sisi kiri (desktop) dan formulir interaksi di sisi kanan untuk menjaga fokus pengguna.

2. **Rancangan Antarmuka Halaman Login (`public/login.php`)**
   - Berfungsi sebagai gerbang autentikasi utama bagi admin/pemilik UMKM untuk mengakses sistem manajemen internal.
   - Dilengkapi kartu logo Berkat Dinasti dengan efek visual *glow*, label penanda sistem manajemen pesanan, indikator status sistem internal yang terverifikasi, formulir input username dan password dengan ikon pemandu, tombol interaktif *toggle* visibilitas kata sandi (*show/hide password*), fitur persistensi sesi (*Remember Me*), serta tautan navigasi menuju pemulihan sandi.

3. **Rancangan Antarmuka Halaman Reset Password (`public/reset-password.php`)**
   - Berfungsi sebagai alur penanganan pemulihan kredensial akun staf/admin secara terkendali.
   - Dilengkapi tombol navigasi kembali (*Kembali ke Login*), formulir pengajuan pemulihan username, serta kartu notifikasi operasional (*Service Level Agreement* 1x24 jam oleh Super Admin) yang memberikan kejelasan proses pemulihan akun bagi pengguna internal.

4. **Modul Pengarah Navigasi Sistem (`index.php`)**
   - Bertindak sebagai pengarah jalur otomatis (*HTTP redirector*) dari direktori utama sistem ke halaman login secara terpusat, konsisten, dan aman.

---

V. Kode Sumber (*Source Code*) Program

Berikut adalah tabel ringkasan berkas kode sumber program komputer yang diimplementasikan pada antarmuka sistem:

| No | Nama Berkas | Direktori (*Path*) | Bahasa Pemrograman | Peran & Deskripsi Fungsional |
|:---:|---|---|---|---|
| 1 | `index.php` | `/index.php` | PHP | Entry Point & Router Redirect ke Antarmuka Login |
| 2 | `login.php` | `/public/login.php` | PHP, HTML5, Tailwind CSS, JS | Antarmuka Autentikasi Pengguna & Branding Bakery |
| 3 | `reset-password.php` | `/public/reset-password.php` | PHP, HTML5, Tailwind CSS | Antarmuka Pemulihan Password Admin Internal |

---

### 5.1 Kode Sumber File: `index.php`
- **Lokasi Berkas**: `/index.php`
- **Bahasa Pemrograman**: PHP
- **Deskripsi Fungsional**: Modul *entry point* navigasi yang secara otomatis mengarahkan akses pengguna ke antarmuka login.

```php
""" + index_code + """
```

---

### 5.2 Kode Sumber File: `public/login.php`
- **Lokasi Berkas**: `/public/login.php`
- **Bahasa Pemrograman**: PHP, HTML5, Tailwind CSS, JavaScript
- **Deskripsi Fungsional**: Modul antarmuka halaman login admin lengkap dengan sistem desain Tailwind CSS, responsive split-screen layout, dan kontrol formulir interaktif.

```php
""" + login_code + """
```

---

### 5.3 Kode Sumber File: `public/reset-password.php`
- **Lokasi Berkas**: `/public/reset-password.php`
- **Bahasa Pemrograman**: PHP, HTML5, Tailwind CSS
- **Deskripsi Fungsional**: Modul antarmuka reset password untuk alur pengajuan pemulihan akun admin internal beserta kartu informasi kebijakan operasional.

```php
""" + reset_code + """
```
"""

    with open(md_path, "w", encoding="utf-8") as f:
        f.write(new_content)
    print(f"Berhasil memperbarui file markdown: {md_path}")

if __name__ == "__main__":
    current_dir = os.path.dirname(os.path.abspath(__file__))
    base_dir = os.path.dirname(current_dir)
    md_file = os.path.join(base_dir, "docs", "DATA", "Kelompok5_HKI Sistem Informasi 40%.md")
    
    # Paths
    out_docx_light = os.path.join(base_dir, "docs", "DATA", "Kelompok5_HKI_Perancangan_UIUX_Berkat_Dinasti.docx")
    out_docx_dark = os.path.join(base_dir, "docs", "DATA", "Kelompok5_HKI_Perancangan_UIUX_Berkat_Dinasti_DarkMode.docx")
    
    docs_root_light = os.path.join(base_dir, "docs", "Kelompok5_HKI_Perancangan_UIUX_Berkat_Dinasti.docx")
    docs_root_dark = os.path.join(base_dir, "docs", "Kelompok5_HKI_Perancangan_UIUX_Berkat_Dinasti_DarkMode.docx")

    # 1. Update Markdown
    update_markdown_file(base_dir, md_file)

    # 2. Generate Primary High-Contrast Modern Light Theme
    generate_full_document(base_dir, out_docx_light, LIGHT_THEME)
    shutil.copy2(out_docx_light, docs_root_light)

    # 3. Generate VS Code Dark Theme Version
    generate_full_document(base_dir, out_docx_dark, DARK_THEME)
    shutil.copy2(out_docx_dark, docs_root_dark)

    print("Semua dokumen DOCX warna-warni berhasil digenerate!")
